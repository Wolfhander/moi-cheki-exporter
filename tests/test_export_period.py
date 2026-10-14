import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from export_period import default_export_years, parse_export_years, filter_by_years
from export import export_receipts_to_excel
from openpyxl import load_workbook


class Settings:
    def get_expense_category_names(self):
        return ["Продукты", "Техника"]


class Database:
    settings = Settings()

    def __init__(self, rows, food=True):
        self.rows = rows
        self.food = food

    def get_receipt_items(self, receipt_id):
        return [{"name": f"Товар-{receipt_id}", "price": 100,
                 "quantity": 1, "total": 100}]

    def get_food_dashboard_data(self):
        return self.rows if self.food else []

    def get_expense_data(self):
        return {"by_category": {"Техника": self.rows}}


def receipt(year, total=100):
    return {"id": year, "receipt_date": f"{year}-12-31T23:59:59+03:00",
            "total_sum": total, "total": total, "price": total, "quantity": 1,
            "seller_name": "Тестовый продавец", "seller_inn": "1234567890",
            "store": "Тестовый магазин", "name": f"Позиция-{year}",
            "category": "Тест", "fiscal_document_id": str(year)}


class ExportPeriodTests(unittest.TestCase):
    def test_defaults_follow_current_calendar_year(self):
        self.assertEqual(default_export_years(datetime(2026, 1, 1)), (2024, 2025, 2026))
        self.assertEqual(default_export_years(datetime(2027, 1, 1)), (2025, 2026, 2027))

    def test_individual_years_ranges_and_all(self):
        self.assertEqual(parse_export_years("2026, 2024–2025, 2024"), (2024, 2025, 2026))
        self.assertEqual(parse_export_years("2024; 2026"), (2024, 2026))
        self.assertIsNone(parse_export_years(" Все "))

    def test_invalid_input(self):
        for value in ("", "2026-2024", "2024,", "abc", "24", "1899", "2024/2026"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_export_years(value)

    def test_dates_and_boundaries(self):
        rows = [{"receipt_date": value} for value in
                ("2023-12-31T23:59:59Z", "2024-01-01", "2024-12-31T23:59:59+03:00",
                 "2025-01-01", "bad", "2024-99-99", None)]
        self.assertEqual(filter_by_years(rows, (2024,)), rows[1:3])
        self.assertIs(filter_by_years(rows, None), rows)

    def test_all_excel_sections_are_filtered(self):
        rows = [receipt(y) for y in (2023, 2024, 2025, 2026)]
        # Ручная покупка и возврат имеют собственные даты и суммы.
        extra = [dict(receipt(2023), name="Старая ручная покупка"),
                 dict(receipt(2024, -30), name="Выбранный возврат", is_ozon_return=True)]
        with tempfile.TemporaryDirectory() as temp:
            path = str(Path(temp) / "report.xlsx")
            export_receipts_to_excel(rows, Database(rows + extra), path, years=(2024, 2026))
            wb = load_workbook(path)
            self.assertEqual(wb["Сводка"]["A3"].value, "Всего чеков: 2")
            self.assertEqual(wb["Сводка"]["A5"].value, "Годы: 2024, 2026")
            self.assertEqual(wb["Все чеки"].max_row, 3)  # заголовок + 2 чека
            self.assertIn("Статистика", wb.sheetnames)
            self.assertIn("Продукты_данные", wb.sheetnames)
            self.assertIn("Техника", wb.sheetnames)
            values = [str(cell.value) for ws in wb for row in ws for cell in row if cell.value is not None]
            text = "\n".join(values)
            self.assertNotIn("Позиция-2023", text)
            self.assertNotIn("Позиция-2025", text)
            self.assertNotIn("Товар-2023", text)
            self.assertNotIn("Товар-2025", text)
            self.assertNotIn("Старая ручная покупка", text)
            self.assertIn("Выбранный возврат", text)
            self.assertIn(-30, [cell.value for row in wb["Продукты_данные"] for cell in row])
            wb.close()
            with zipfile.ZipFile(path) as archive:
                charts = "".join(archive.read(n).decode() for n in archive.namelist()
                                 if n.startswith("xl/charts/chart") and n.endswith(".xml"))
                self.assertNotIn("2023", charts)
                self.assertNotIn("2025", charts)

    def test_no_food_does_not_expose_template_data(self):
        with tempfile.TemporaryDirectory() as temp:
            path = str(Path(temp) / "empty-food.xlsx")
            export_receipts_to_excel([receipt(2024)], Database([receipt(2023)]), path, years=(2024,))
            wb = load_workbook(path)
            self.assertNotIn("Продукты", wb.sheetnames)
            wb.close()

    def test_all_years_and_manual_only_export(self):
        with tempfile.TemporaryDirectory() as temp:
            path = str(Path(temp) / "all.xlsx")
            rows = [receipt(2023), receipt(2026)]
            export_receipts_to_excel(rows, Database(rows), path, years=None)
            wb = load_workbook(path)
            self.assertEqual(wb["Сводка"]["A3"].value, "Всего чеков: 2")
            wb.close()
            export_receipts_to_excel([], Database([receipt(2024)]), path, years=(2024,))
            wb = load_workbook(path)
            self.assertIn("Продукты_данные", wb.sheetnames)
            wb.close()


if __name__ == "__main__":
    unittest.main()
